from docx import Document
from docx.shared import Pt

content_title = "SaaS Migration Technical Runbook\n"

sections = [
    ("Executive Summary", [
        "Goal: Convert the existing training site into a multi-tenant SaaS so different training institutes can create and manage their own students, courses, and billing.",
        "Outcome: A self-serve platform with tenant isolation, recurring billing, and admin tools for support and reporting."
    ]),

    ("Assumptions", [
        "App is Laravel-based (PHP) with MySQL and S3-style storage.",
        "Current payments use Paystack for one-off registration; we will add Stripe (recommended) for subscriptions or keep Paystack recurring where regionally required."
    ]),

    ("Phases & Steps", [
        "Phase 0 — Discovery (2–4 days): inventory models/tables/controllers; produce table→model→controller mapping.",
        "Phase 1 — Tenancy scaffold (shared-schema) (2–3 weeks): create Tenant model + migrations; add nullable tenant_id to core tables; implement TenantAware trait with global Eloquent scope; ResolveTenant middleware (subdomain/org slug); CLI helpers; update controllers to use tenant scoping; add tests.",
        "Phase 2 — Billing & subscriptions (1–2 weeks): install Stripe + Laravel Cashier; add billing fields to tenants; implement checkout/subscription, webhook handlers; billing UI or use Stripe Billing Portal; migrate or coexist with Paystack for legacy flows.",
        "Phase 3 — Onboarding & admin UI (2–3 weeks): public signup → create tenant + admin user; tenant dashboard (settings, team invites, billing, usage); quotas enforcement and notifications.",
        "Phase 4 — Data migration & rollout (1–2 weeks): create default tenant; backup DB; migration script to set tenant_id on existing rows; test in staging; make tenant_id non-nullable after verification.",
        "Phase 5 — Testing, CI/CD & staging (1–2 weeks): unit, integration, E2E tests; GitHub Actions for CI (phpunit, static analysis); staging deploy and smoke tests.",
        "Phase 6 — Operations & monitoring (ongoing): backups, Sentry, APM, centralized logs, uptime monitoring, runbooks.",
        "Phase 7 — Enterprise features (later): SAML/SCIM SSO, per-tenant DB option, custom domains, SLA and dedicated support."
    ]),

    ("Key Data Model Changes (concrete)", [
        "Add `tenants` table: id, name, slug, stripe_customer_id, stripe_subscription_id, status, settings (json), timestamps.",
        "Add `tenant_id` to core tables: lms_courses, lms_modules, lms_materials, lms_students, lms_enrollments, lms_tracks, lms_scheduled_classes, lms_certificates, payments, training_registrations.",
        "Create `Tenant` model and `TenantAware` trait (global scope) and `ResolveTenant` middleware.",
    ]),

    ("Billing Strategy (simple)", [
        "Use Stripe + Laravel Cashier for subscriptions, trials, invoices, webhooks. Keep Paystack for legacy one-offs or migrate to Stripe.",
        "Store provider IDs on `tenants` (stripe_customer_id, stripe_subscription_id) and update billing_status via webhooks.",
        "Secure webhook verification and do not store card numbers."
    ]),

    ("Migration Plan (safe)", [
        "1. Add nullable tenant_id columns and create `default` tenant.",
        "2. Back up DB and run migration script to set existing rows' tenant_id to the default tenant.",
        "3. Deploy middleware and TenantAware in read-only mode; validate in staging.",
        "4. After verification, make tenant_id non-nullable and enforce strict scoping."
    ]),

    ("Testing & QA (must-have)", [
        "Unit tests for TenantAware scopes and billing logic.",
        "Integration tests for signup -> tenant creation -> subscription -> create course -> enroll flows.",
        "E2E tests for full happy paths using Cypress or Playwright."
    ]),

    ("Security & Compliance (simple checklist)", [
        "HTTPS everywhere, HSTS, CSP.",
        "Webhook and API signature verification.",
        "GDPR endpoints: data export and deletion for tenants.",
        "Secrets in environment/secret manager (no .env in repo)."
    ]),

    ("CI/CD & Infra", [
        "GitHub Actions: run PHP CS Fixer, PHPStan, phpunit on PRs.",
        "Environments: dev, staging, prod. Use managed platform (Render, DigitalOcean App Platform, or Laravel Vapor).",
        "Storage: S3-compatible for media. Backups: nightly DB dumps + retention."
    ]),

    ("Estimated Team & Timeline", [
        "Team: 1–2 backend devs (Laravel), 1 frontend/designer part-time, 1 devops part-time, 1 QA part-time.",
        "MVP timeline: 6–10 weeks total depending on team size and parallel work."
    ]),

    ("Immediate Next Steps", [
        "Approve 2–4 day discovery audit to produce a detailed file-by-file list and firm cost estimate.",
        "Decide on primary payment provider (Stripe recommended) or confirm local provider requirement.",
        "If approved, run automated repo audit and scaffold Tenant model, middleware, and a migration to add tenant_id columns."
    ])
]


def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    if bold:
        run.bold = True


doc = Document()
doc.add_heading('SaaS Migration Technical Runbook', level=1)

doc.add_paragraph('This document summarizes the technical plan to convert the existing training site into a multi-tenant SaaS platform where each training institute can sign up and manage its own students, courses, and billing.')

for title, bullets in sections:
    doc.add_heading(title, level=2)
    for b in bullets:
        p = doc.add_paragraph(style='List Bullet')
        run = p.add_run(b)
        run.font.size = Pt(11)

# Footer note

doc.add_paragraph('\nGenerated by an automated script. Review content and adjust timelines based on team availability.')

out_path = 'docs/saas_plan.docx'
try:
    doc.save(out_path)
    print(f'WROTE:{out_path}')
except Exception as e:
    print('ERROR:', e)
